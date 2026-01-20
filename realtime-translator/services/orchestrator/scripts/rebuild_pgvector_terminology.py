from __future__ import annotations

import os
from pathlib import Path
from typing import Dict, List, Tuple

import psycopg
from psycopg import sql
from openpyxl import load_workbook
from pgvector.psycopg import register_vector
from sentence_transformers import SentenceTransformer

# Connection string to your Docker Postgres
DSN = "postgresql://postgres:oebb@127.0.0.1:5432/terminology"

# Must match your table definition vector(384)
MODEL_NAME = "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"

# Resource control
BATCH_SIZE = 25

DEFAULT_MODE = "force"
DEFAULT_PRIORITY = 100

# Paths
SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parents[2]  # ...\realtime-translator
TERMINOLOGY_DIR = PROJECT_ROOT / "rag" / "terminology"

XLSX_DE = TERMINOLOGY_DIR / "railway_terminology.de.xlsx"
XLSX_EN = TERMINOLOGY_DIR / "railway_terminology.en.xlsx"


def _s(x) -> str:
    return "" if x is None else str(x).strip()


def _pick_header(headers: List[str], candidates: List[str]) -> str:
    for c in candidates:
        if c in headers:
            return c
    raise SystemExit(f"Missing header. Need one of {candidates}. Found (first 40): {headers[:40]}")


def _read_rows(
    xlsx: Path,
    term_candidates: List[str],
    def_candidates: List[str],
    lang_candidates: List[str],
    lang_value: str,
) -> List[Tuple[str, str]]:
    wb = load_workbook(filename=xlsx, read_only=True, data_only=True)
    ws = wb.active

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [_s(h) for h in header_row]

    term_h = _pick_header(headers, term_candidates)
    def_h = _pick_header(headers, def_candidates)
    lang_h = _pick_header(headers, lang_candidates)

    term_i = headers.index(term_h)
    def_i = headers.index(def_h)
    lang_i = headers.index(lang_h)

    out: List[Tuple[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        term = _s(row[term_i])
        definition = _s(row[def_i])
        lang = _s(row[lang_i]).lower()

        if not term:
            continue

        # allow "de-DE", "en-US", etc.
        if not lang.startswith(lang_value):
            continue

        out.append((term, definition))

    wb.close()
    return out


def ensure_schema(conn: psycopg.Connection) -> None:
    with conn.cursor() as cur:
        cur.execute("CREATE EXTENSION IF NOT EXISTS vector;")

        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS public.glossary_terms (
              id bigserial PRIMARY KEY,
              src_lang text NOT NULL,
              tgt_lang text NOT NULL,
              src_term text NOT NULL,
              tgt_term text NOT NULL,
              mode text NOT NULL,
              priority int NOT NULL DEFAULT 0,
              meta jsonb NULL,
              embedding vector(384) NOT NULL
            );
            """
        )

        # correct unique constraint
        cur.execute(
            """
            DO $$
            BEGIN
              IF NOT EXISTS (
                SELECT 1 FROM pg_constraint
                WHERE conname = 'glossary_terms_uniq'
                  AND conrelid = 'public.glossary_terms'::regclass
              ) THEN
                ALTER TABLE public.glossary_terms
                ADD CONSTRAINT glossary_terms_uniq
                UNIQUE (src_lang, tgt_lang, src_term, tgt_term, mode);
              END IF;
            END $$;
            """
        )

        cur.execute(
            "CREATE INDEX IF NOT EXISTS glossary_terms_lang_idx ON public.glossary_terms (src_lang, tgt_lang);"
        )

    conn.commit()


def recreate_vector_index(conn: psycopg.Connection) -> None:
    with conn.cursor() as cur:
        cur.execute("DROP INDEX IF EXISTS public.glossary_terms_embedding_idx;")
        conn.commit()

        # Try HNSW first, fallback to IVFFLAT
        try:
            cur.execute(
                """
                CREATE INDEX glossary_terms_embedding_idx
                ON public.glossary_terms
                USING hnsw (embedding vector_cosine_ops);
                """
            )
            conn.commit()
            return
        except Exception:
            conn.rollback()

        cur.execute(
            """
            CREATE INDEX glossary_terms_embedding_idx
            ON public.glossary_terms
            USING ivfflat (embedding vector_cosine_ops)
            WITH (lists = 50);
            """
        )
    conn.commit()


def main() -> None:
    os.environ["TOKENIZERS_PARALLELISM"] = "false"
    os.environ["OMP_NUM_THREADS"] = "1"
    os.environ["MKL_NUM_THREADS"] = "1"

    if not XLSX_DE.exists():
        raise SystemExit(f"DE file not found: {XLSX_DE}")
    if not XLSX_EN.exists():
        raise SystemExit(f"EN file not found: {XLSX_EN}")

    de_rows = _read_rows(
        XLSX_DE,
        term_candidates=["Begriff", "term", "Term"],
        def_candidates=["Definition", "definition", "Definition (DE)"],
        lang_candidates=["Sprache", "language", "Lang"],
        lang_value="de",
    )
    en_rows = _read_rows(
        XLSX_EN,
        term_candidates=["term", "Term", "Begriff"],
        def_candidates=["definition", "Definition", "Definition (EN)"],
        lang_candidates=["language", "Sprache", "Lang"],
        lang_value="en",
    )

    print("DE rows:", len(de_rows))
    print("EN rows:", len(en_rows))

    if len(de_rows) != len(en_rows):
        raise SystemExit(
            f"Row mismatch: de={len(de_rows)} en={len(en_rows)}. "
            "If the files are not row-aligned, you need a shared ID column."
        )

    model = SentenceTransformer(MODEL_NAME)

    with psycopg.connect(DSN) as conn:
        register_vector(conn)
        ensure_schema(conn)

        with conn.cursor() as cur:
            cur.execute("TRUNCATE TABLE public.glossary_terms;")
        conn.commit()

        total_inserted = 0

        # Build and insert in small batches
        for i in range(0, len(de_rows), BATCH_SIZE):
            chunk_de = de_rows[i : i + BATCH_SIZE]
            chunk_en = en_rows[i : i + BATCH_SIZE]

            texts: List[str] = []
            pairs: List[Tuple[str, str, str, str]] = []

            for (de_term, de_def), (en_term, en_def) in zip(chunk_de, chunk_en):
                texts.append("\n".join([en_term, en_def, de_term, de_def]).strip())
                pairs.append((en_term, en_def, de_term, de_def))

            embs = model.encode(texts, normalize_embeddings=True).tolist()

            rows = []
            for (en_term, en_def, de_term, de_def), emb in zip(pairs, embs):
                meta = {"de_def": de_def, "en_def": en_def}

                rows.append(
                    ("en", "de", en_term, de_term, DEFAULT_MODE, DEFAULT_PRIORITY, meta, emb)
                )
                rows.append(
                    ("de", "en", de_term, en_term, DEFAULT_MODE, DEFAULT_PRIORITY, meta, emb)
                )

            with conn.cursor() as cur:
                cur.executemany(
                    """
                    INSERT INTO public.glossary_terms
                      (src_lang, tgt_lang, src_term, tgt_term, mode, priority, meta, embedding)
                    VALUES
                      (%s,%s,%s,%s,%s,%s,%s,%s)
                    ON CONFLICT (src_lang, tgt_lang, src_term, tgt_term, mode)
                    DO UPDATE SET
                      priority = EXCLUDED.priority,
                      meta = EXCLUDED.meta,
                      embedding = EXCLUDED.embedding
                    """,
                    rows,
                )
            conn.commit()
            total_inserted += len(rows)
            print("Inserted so far:", total_inserted)

        recreate_vector_index(conn)

        with conn.cursor() as cur:
            cur.execute("ANALYZE public.glossary_terms;")
            cur.execute("SELECT COUNT(*) FROM public.glossary_terms;")
            count = cur.fetchone()[0]
        conn.commit()

    print("Done. DB row count:", count)
    print("Terminology dir:", TERMINOLOGY_DIR)


if __name__ == "__main__":
    main()
